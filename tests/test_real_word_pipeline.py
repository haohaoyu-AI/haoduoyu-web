"""Offline integration test using a conventional third-party generated DOCX.

The existing workflow tests intentionally use article_hub.docxio.write_document
for compact fixtures. This test exercises the same import and layout paths with
an independently generated OOXML Word file and real image files.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from PIL import Image, ImageDraw

from article_hub import config, image_review, library, topic_packages
from article_hub.docxio import read_document


class RealWordPipelineTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="haoduoyu-real-word-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.unformatted = self.root / "未排版文章"
        self.formatted = self.root / "已排版文章"
        self.data = self.root / "data"
        self.preview = self.root / "预览"
        for path in (self.unformatted, self.formatted, self.data, self.preview):
            path.mkdir()
        self._patches = []
        for name, value in (("UNFORMATTED", self.unformatted),
                            ("FORMATTED", self.formatted),
                            ("DATA", self.data), ("PREVIEW", self.preview)):
            p = patch.object(config, name, value)
            p.start()
            self._patches.append(p)
        old_ledger = library.LEDGER
        library.LEDGER = self.data / "库.json"
        self.addCleanup(setattr, library, "LEDGER", old_ledger)
        self.addCleanup(self._stop_patches)

    def _stop_patches(self):
        for p in reversed(self._patches):
            p.stop()

    def _build_topic(self):
        topic_id = "TP-20261001-a1b2c3d4"
        title = "测试样本 城市公交站点改造"
        source = "城市公交站点改造"
        package = topic_packages.build(
            topic_id=topic_id,
            title=title,
            original_topic=source,
            keyword_source=source,
            source="offline synthetic fixture",
            generated={
                "selected_title": title,
                "angle": "从无障碍通行和候车空间观察站点改造",
                "outline": [],
                "prompt": "请依据测试素材撰写可核验的公交站点改造说明，不添加真实人名、地点或数据。",
                "image_plan": [{
                    "position_id": "body-1", "order": 0,
                    "query": source, "keyword_type": "core",
                }],
            },
        )
        folder = topic_packages.commit(package)
        topic_packages.mark_imported(topic_id)
        return topic_id, title, folder

    def _make_image(self, path: Path, label: str, color: tuple[int, int, int]):
        image = Image.new("RGB", (720, 480), color)
        draw = ImageDraw.Draw(image)
        draw.rectangle((50, 60, 670, 300), fill=(232, 240, 230))
        draw.rectangle((130, 170, 590, 300), fill=(60, 100, 140))
        draw.ellipse((185, 275, 260, 350), fill=(30, 30, 30))
        draw.ellipse((455, 275, 530, 350), fill=(30, 30, 30))
        draw.text((70, 385), label, fill=(20, 20, 20))
        image.save(path, format="PNG")

    def _make_external_word(self, path: Path, title: str):
        doc = Document()
        doc.add_paragraph(title, style="Title")
        doc.add_heading("改造缘起", level=2)
        doc.add_paragraph("合成测试正文：本段描述候车区域与步行通道的关系。")
        doc.add_heading("设计要点", level=2)
        doc.add_paragraph("合成测试正文：本段用于检查图片插入位置和段落顺序。")
        doc.save(path)

    def _preserve_artifacts(self, topic_id, folder, final, rollback_record):
        destination = os.environ.get("HAODOUYU_REAL_WORD_ARTIFACTS")
        if not destination:
            return
        out = Path(destination)
        for name in ("input", "output", "rollback.json"):
            existing = out / name
            if existing.is_dir() and not existing.is_symlink():
                shutil.rmtree(existing)
            else:
                existing.unlink(missing_ok=True)
        (out / "input").mkdir(parents=True)
        (out / "output").mkdir()
        source = topic_packages.imported_word(image_review.topic_path("成稿:" + topic_id))
        shutil.copy2(source, out / "input" / source.name)
        for name in ("fixture-approved.png", "fixture-rejected.png", "fixture-pending.png"):
            shutil.copy2(folder / name, out / "input" / name)
        shutil.copy2(final, out / "output" / final.name)
        shutil.copy2(final.parent / "配图清单.json", out / "output" / "配图清单.json")
        shutil.copytree(final.parent / "配图", out / "output" / "配图")
        shutil.copytree(final.parent / "原始素材", out / "output" / "原始素材")
        (out / "rollback.json").write_text(
            json.dumps(rollback_record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def test_external_docx_import_manual_review_layout_archive_and_rollback(self):
        topic_id, title, folder = self._build_topic()
        fixtures = {
            "fixture-approved.png": ("TEST APPROVED IMAGE", (95, 155, 195)),
            "fixture-rejected.png": ("TEST REJECTED IMAGE", (195, 115, 95)),
            "fixture-pending.png": ("TEST PENDING IMAGE", (125, 175, 110)),
        }
        for name, (label, color) in fixtures.items():
            self._make_image(folder / name, label, color)
        names = list(fixtures)
        image_review.save_candidate_review(topic_id, names)
        image_review.save_review(topic_id, {
            names[0]: "approved", names[1]: "rejected", names[2]: "pending",
        })

        export = image_review.topic_export_folder(topic_id)
        self.assertIsNotNone(export)
        self._make_external_word(export / "synthetic-source.docx", title)
        original_docx = (export / "synthetic-source.docx").read_bytes()
        imported = topic_packages.import_word_from_export(
            topic_id, filename="synthetic-source.docx")
        self.assertEqual(imported["filename"], "synthetic-source.docx")
        self.assertEqual(topic_packages.imported_word(folder).read_bytes(), original_docx)
        self.assertEqual(topic_packages.import_word_from_export(topic_id)["already_imported"], True)

        review = image_review.load_review(folder)
        self.assertEqual(review["states"], {
            names[0]: "approved", names[1]: "rejected", names[2]: "pending",
        })
        package_files_before = {
            path.name: path.read_bytes() for path in folder.iterdir()
            if path.is_file() and path.name != image_review.REVIEW_FILE
        }
        image_review.reformat(topic_id)

        finished = topic_packages.formatted_folder_for(topic_id)
        self.assertIsNotNone(finished)
        final = finished / (title + ".docx")
        self.assertTrue(final.is_file())
        self.assertEqual(read_document(final).title, title)
        reopened = Document(final)
        paragraphs = reopened.paragraphs
        text = [paragraph.text for paragraph in paragraphs]
        self.assertEqual([value for value in text if value], [
            title, "改造缘起", "合成测试正文：本段描述候车区域与步行通道的关系。",
            "设计要点", "合成测试正文：本段用于检查图片插入位置和段落顺序。",
        ])
        self.assertEqual(len(reopened.inline_shapes), 1)
        image_paragraphs = [index for index, paragraph in enumerate(paragraphs)
                            if paragraph._p.xpath(".//w:drawing")]
        self.assertEqual(len(image_paragraphs), 1)
        image_index = image_paragraphs[0]
        self.assertGreater(image_index, 1, "新版正文配图不得被挪到标题下作为封面")
        before_image = next(paragraph.text for paragraph in reversed(paragraphs[:image_index])
                            if paragraph.text.strip())
        after_image = next(paragraph.text for paragraph in paragraphs[image_index + 1:]
                           if paragraph.text.strip())
        self.assertEqual(before_image, "合成测试正文：本段描述候车区域与步行通道的关系。")
        self.assertEqual(after_image, "设计要点")
        manifest = json.loads((finished / "配图清单.json").read_text(encoding="utf-8"))
        used = [row for row in manifest["配图"] if row.get("角色") != "未使用"]
        self.assertEqual([row["原始文件名"] for row in used], [names[0]])
        output_image = finished / "配图" / used[0]["文件"]
        self.assertTrue(output_image.is_file())
        with Image.open(output_image) as image:
            image.verify()
        self.assertNotIn(names[1], [row["原始文件名"] for row in used])
        self.assertNotIn(names[2], [row["原始文件名"] for row in used])
        archived = finished / "原始素材"
        self.assertTrue(archived.is_dir())
        for name, data in package_files_before.items():
            self.assertEqual((archived / name).read_bytes(), data, name)
        self.assertEqual((archived / "synthetic-source.docx").read_bytes(), original_docx)
        archived_review = image_review.load_review(archived)
        self.assertEqual(archived_review["states"], review["states"])
        self.assertFalse(archived_review["needs_layout"])

        snapshot = {
            "final_docx_sha256_before_failure": hashlib.sha256(final.read_bytes()).hexdigest(),
            "manifest_sha256_before_failure": hashlib.sha256(
                (finished / "配图清单.json").read_bytes()).hexdigest(),
            "approved_image_sha256_before_failure": hashlib.sha256(
                output_image.read_bytes()).hexdigest(),
        }
        with patch.object(library, "refresh_ledger", side_effect=OSError("synthetic ledger commit failure")):
            with self.assertRaisesRegex(OSError, "synthetic ledger commit failure"):
                image_review.reformat("成稿:" + topic_id)
        self.assertEqual(hashlib.sha256(final.read_bytes()).hexdigest(),
                         snapshot["final_docx_sha256_before_failure"])
        self.assertEqual(hashlib.sha256((finished / "配图清单.json").read_bytes()).hexdigest(),
                         snapshot["manifest_sha256_before_failure"])
        self.assertEqual(hashlib.sha256(output_image.read_bytes()).hexdigest(),
                         snapshot["approved_image_sha256_before_failure"])
        reopened_after_failure = Document(final)
        self.assertEqual(reopened_after_failure.paragraphs[0].text, title)
        self.assertEqual(len(reopened_after_failure.inline_shapes), 1)
        self.assertTrue(list((self.data / "重新排版备份").glob("*/未提交成稿")))
        snapshot["after_failure"] = {
            "docx_sha256": hashlib.sha256(final.read_bytes()).hexdigest(),
            "manifest_sha256": hashlib.sha256((finished / "配图清单.json").read_bytes()).hexdigest(),
            "approved_image_sha256": hashlib.sha256(
                output_image.read_bytes()).hexdigest(),
            "reopen_ok": True,
            "inline_images": len(reopened_after_failure.inline_shapes),
        }
        image_review.reformat("成稿:" + topic_id)
        reopened_after_retry = Document(final)
        self.assertEqual(reopened_after_retry.paragraphs[0].text, title)
        self.assertEqual(len(reopened_after_retry.inline_shapes), 1)
        self.assertEqual(image_review.load_review(archived)["states"], review["states"])
        snapshot["retry_after_failure"] = {
            "reopen_ok": True,
            "inline_images": len(reopened_after_retry.inline_shapes),
            "review_states_preserved": image_review.load_review(archived)["states"] == review["states"],
        }
        self._preserve_artifacts(topic_id, archived, final, snapshot)


if __name__ == "__main__":
    unittest.main()
