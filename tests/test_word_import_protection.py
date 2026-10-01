"""Offline import-safety checks for corrupt and replacement DOCX files."""

from __future__ import annotations

import hashlib
import io
import json
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from PIL import Image

from article_hub import config, console, image_review, library, topic_packages


class WordImportProtectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="haoduoyu-word-import-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        for name in ("UNFORMATTED", "FORMATTED", "DATA", "PREVIEW"):
            value = self.root / name
            value.mkdir()
            p = patch.object(config, name, value)
            p.start()
            self.addCleanup(p.stop)
        old_ledger = library.LEDGER
        library.LEDGER = self.root / "DATA" / "库.json"
        self.addCleanup(setattr, library, "LEDGER", old_ledger)

    def _package_with_approved_image(self):
        topic_id = "TP-20261001-11112222"
        title = "测试样本 公交站点"
        package = topic_packages.build(
            topic_id=topic_id,
            title=title,
            original_topic=title,
            keyword_source=title,
            source="offline fixture",
            generated={
                "selected_title": title,
                "angle": "观察测试样本中的公交站点设计",
                "outline": [],
                "prompt": "请只使用无个人信息的临时样本，逐项说明公交站点设计测试内容。",
                "image_plan": [{"position_id": "body-1", "order": 0,
                                "query": title, "keyword_type": "core"}],
            },
        )
        folder = topic_packages.commit(package)
        topic_packages.mark_imported(topic_id)
        image_path = folder / "fixture-approved.png"
        Image.new("RGB", (640, 420), (34, 129, 177)).save(image_path, format="PNG")
        image_review.save_candidate_review(topic_id, [image_path.name])
        image_review.save_review(topic_id, {image_path.name: "approved"})
        return topic_id, folder, image_path.name

    @staticmethod
    def _word_bytes(title: str, body: str) -> bytes:
        output = io.BytesIO()
        document = Document()
        document.add_paragraph(title, style="Title")
        document.add_heading("测试正文", level=2)
        document.add_paragraph(body)
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _post_word(topic_id: str, filename: str, data: bytes):
        handler = object.__new__(console.Handler)
        handler.path = ("/api/topic-word?topic=" + topic_id + "&name=" + filename)
        handler.headers = {"Content-Length": str(len(data))}
        handler.rfile = io.BytesIO(data)
        responses = []
        handler._json = lambda payload, status=200: responses.append((status, payload))
        handler._handle_post()
        return responses[0]

    @staticmethod
    def _assert_no_partial_import(test, folder, filename):
        test.assertIsNone(topic_packages.imported_word(folder))
        test.assertNotIn("word", topic_packages.load(folder))
        test.assertFalse((folder / filename).exists())
        test.assertEqual(list(folder.glob(".*.docx")), [])

    @staticmethod
    def _preserve_artifacts(folder, filename, image_name, inputs, attempts):
        destination = os.environ.get("HAODOUYU_WORD_IMPORT_ARTIFACTS")
        if not destination:
            return
        out = Path(destination)
        for name in ("input", "output", "attempts.json", "SHA256SUMS"):
            target = out / name
            if target.is_dir() and not target.is_symlink():
                shutil.rmtree(target)
            else:
                target.unlink(missing_ok=True)
        (out / "input").mkdir(parents=True)
        (out / "output").mkdir()
        for name, content in inputs.items():
            (out / "input" / name).write_bytes(content)
        shutil.copy2(folder / filename, out / "output" / filename)
        shutil.copy2(folder / image_name, out / "output" / image_name)
        shutil.copy2(folder / topic_packages.PACKAGE_FILE, out / "output" / topic_packages.PACKAGE_FILE)
        shutil.copy2(folder / image_review.REVIEW_FILE, out / "output" / image_review.REVIEW_FILE)
        (out / "attempts.json").write_text(
            json.dumps(attempts, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def test_corrupt_files_retry_and_same_name_replacement_protection(self):
        topic_id, folder, image_name = self._package_with_approved_image()
        package_before = (folder / topic_packages.PACKAGE_FILE).read_bytes()
        review_before = (folder / image_review.REVIEW_FILE).read_bytes()
        filename = "synthetic-source.docx"
        package_before_hash = hashlib.sha256(package_before).hexdigest()
        review_before_hash = hashlib.sha256(review_before).hexdigest()

        # A valid ZIP container without the required main OOXML document.
        missing_document = io.BytesIO()
        with zipfile.ZipFile(missing_document, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/styles.xml", "<styles/>")
        bad_zip_parts = self._word_bytes("测试标题", "截断前的合成正文")
        truncated = bad_zip_parts[:-48]

        attempts = []
        for label, corrupt in (("missing document.xml", missing_document.getvalue()),
                               ("truncated ZIP", truncated)):
            with self.subTest(corrupt=label):
                status, body = self._post_word(topic_id, filename, corrupt)
                attempts.append({"case": label, "status": status, "error": body.get("error")})
                self.assertEqual(status, 400)
                self.assertIn("error", body)
                self._assert_no_partial_import(self, folder, filename)
                self.assertEqual((folder / topic_packages.PACKAGE_FILE).read_bytes(), package_before)
                self.assertEqual((folder / image_review.REVIEW_FILE).read_bytes(), review_before)
                self.assertEqual(image_review.load_review(folder)["states"], {image_name: "approved"})
                attempts[-1].update({
                    "no_partial_import": topic_packages.imported_word(folder) is None,
                    "package_hash_unchanged": hashlib.sha256(
                        (folder / topic_packages.PACKAGE_FILE).read_bytes()).hexdigest() == package_before_hash,
                    "review_hash_unchanged": hashlib.sha256(
                        (folder / image_review.REVIEW_FILE).read_bytes()).hexdigest() == review_before_hash,
                    "approved_state_unchanged": image_review.load_review(folder)["states"].get(image_name) == "approved",
                })

        # A valid retry after both failures binds the selected file.
        original = self._word_bytes("测试样本 公交站点", "第一份有效合成稿的正文。")
        status, body = self._post_word(topic_id, filename, original)
        attempts.append({"case": "valid retry", "status": status, "error": body.get("error")})
        self.assertEqual(status, 200, body)
        bound = topic_packages.imported_word(folder)
        self.assertIsNotNone(bound)
        self.assertEqual(bound.read_bytes(), original)
        bound_hash = hashlib.sha256(bound.read_bytes()).hexdigest()
        package_after_bind = (folder / topic_packages.PACKAGE_FILE).read_bytes()
        package_after_bind_hash = hashlib.sha256(package_after_bind).hexdigest()

        # Same name, different content must not silently replace the bound source.
        replacement = self._word_bytes("测试样本 公交站点", "第二份不同内容的合成稿，未经确认。")
        status, body = self._post_word(topic_id, filename, replacement)
        attempts.append({"case": "different content, same bound filename", "status": status,
                         "error": body.get("error")})
        self.assertEqual(status, 400)
        self.assertIn("已经导入过 Word", body["error"])
        self.assertEqual(hashlib.sha256(bound.read_bytes()).hexdigest(), bound_hash)
        self.assertEqual((folder / topic_packages.PACKAGE_FILE).read_bytes(), package_after_bind)
        self.assertEqual((folder / image_review.REVIEW_FILE).read_bytes(), review_before)
        self.assertEqual(image_review.load_review(folder)["states"], {image_name: "approved"})
        attempts[-1].update({
            "bound_hash_unchanged": hashlib.sha256(bound.read_bytes()).hexdigest() == bound_hash,
            "package_hash_unchanged": hashlib.sha256(
                (folder / topic_packages.PACKAGE_FILE).read_bytes()).hexdigest() == package_after_bind_hash,
            "review_hash_unchanged": hashlib.sha256(
                (folder / image_review.REVIEW_FILE).read_bytes()).hexdigest() == review_before_hash,
        })

        # The title folder may hold another Word. Selecting it after binding is
        # explicitly rejected by current code; there is no replacement-confirm
        # flag or UI flow, so the test preserves that existing protection.
        export = image_review.topic_export_folder(topic_id)
        self.assertIsNotNone(export)
        (export / filename).write_bytes(replacement)
        with self.assertRaisesRegex(ValueError, "已绑定另一份 Word"):
            topic_packages.import_word_from_export(topic_id, filename=filename)
        attempts.append({"case": "changed same-name file in title folder", "status": "rejected",
                         "error": "已绑定另一份 Word"})
        self.assertEqual(hashlib.sha256(bound.read_bytes()).hexdigest(), bound_hash)
        self.assertEqual((folder / topic_packages.PACKAGE_FILE).read_bytes(), package_after_bind)
        self.assertEqual((folder / image_review.REVIEW_FILE).read_bytes(), review_before)
        self.assertEqual(image_review.load_review(folder)["states"], {image_name: "approved"})
        attempts[-1].update({
            "bound_hash_unchanged": hashlib.sha256(bound.read_bytes()).hexdigest() == bound_hash,
            "package_hash_unchanged": hashlib.sha256(
                (folder / topic_packages.PACKAGE_FILE).read_bytes()).hexdigest() == package_after_bind_hash,
            "review_hash_unchanged": hashlib.sha256(
                (folder / image_review.REVIEW_FILE).read_bytes()).hexdigest() == review_before_hash,
        })

        # Choosing a specific second title-folder Word is only permitted before
        # a source is bound; the existing multiple-file test covers that path.
        # After binding, choosing another candidate must not silently replace it.
        alternate = self._word_bytes("测试样本 公交站点", "同题第二份候选 Word。")
        (export / "alternative.docx").write_bytes(alternate)
        with self.assertRaisesRegex(ValueError, "已绑定另一份 Word"):
            topic_packages.import_word_from_export(topic_id, filename="alternative.docx")
        attempts.append({"case": "second title-folder Word after binding", "status": "rejected",
                         "error": "已绑定另一份 Word"})
        self.assertEqual(hashlib.sha256(bound.read_bytes()).hexdigest(), bound_hash)
        self.assertEqual(image_review.load_review(folder)["states"], {image_name: "approved"})
        attempts[-1].update({
            "bound_hash_unchanged": hashlib.sha256(bound.read_bytes()).hexdigest() == bound_hash,
            "package_hash_unchanged": hashlib.sha256(
                (folder / topic_packages.PACKAGE_FILE).read_bytes()).hexdigest() == package_after_bind_hash,
            "review_hash_unchanged": hashlib.sha256(
                (folder / image_review.REVIEW_FILE).read_bytes()).hexdigest() == review_before_hash,
        })
        self._preserve_artifacts(folder, filename, image_name, {
            "missing-document.docx": missing_document.getvalue(),
            "truncated.docx": truncated,
            "synthetic-source.docx": original,
            "different-content-same-name.docx": replacement,
            "alternative.docx": alternate,
        }, {"attempts": attempts, "bound_sha256": bound_hash,
            "package_sha256_after_binding": package_after_bind_hash,
            "review_sha256_before_and_after": review_before_hash,
            "review_states": image_review.load_review(folder)["states"]})


if __name__ == "__main__":
    unittest.main()
