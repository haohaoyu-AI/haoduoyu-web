"""Build a standards-based upload copy; leave the source and queue identity intact."""
import io
import posixpath
import tempfile
import urllib.parse
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches
from .articles import load_article
from PIL import Image

_RELATIONSHIP_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def pixels(raw):
    with Image.open(io.BytesIO(raw)) as image:
        return image.size, image.convert('RGBA').tobytes()


def _source_part_for_relationships(name):
    if name == "_rels/.rels":
        return ""
    if "/_rels/" not in name or not name.endswith(".rels"):
        return None
    directory, part = name.rsplit("/_rels/", 1)
    return posixpath.join(directory, part[:-5])


def _missing_image_relationships(source):
    """Find dangling internal image links that make python-docx reject a DOCX."""
    missing = {}
    with zipfile.ZipFile(source) as archive:
        names = set(archive.namelist())
        for rels_name in names:
            source_part = _source_part_for_relationships(rels_name)
            if source_part is None:
                continue
            try:
                root = ET.fromstring(archive.read(rels_name))
            except ET.ParseError:
                continue
            for relation in root:
                if (not relation.get("Type", "").endswith("/image")
                        or relation.get("TargetMode", "").lower() == "external"):
                    continue
                target = urllib.parse.unquote(relation.get("Target", "")).replace("\\", "/")
                if not target:
                    continue
                member = (posixpath.normpath(target.lstrip("/")) if target.startswith("/")
                          else posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target)))
                if member not in names:
                    missing.setdefault(rels_name, set()).add(relation.get("Id", ""))
    return {name: ids - {""} for name, ids in missing.items() if ids - {""}}


def _write_without_missing_images(source, target, missing):
    """Copy a DOCX while dropping only relationships to image parts that do not exist."""
    ET.register_namespace("", _RELATIONSHIP_NS)
    with zipfile.ZipFile(source, "r") as original, zipfile.ZipFile(target, "w") as repaired:
        for info in original.infolist():
            data = original.read(info.filename)
            invalid_ids = missing.get(info.filename)
            if invalid_ids:
                root = ET.fromstring(data)
                for relation in list(root):
                    if relation.get("Id") in invalid_ids:
                        root.remove(relation)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            repaired.writestr(info, data)


def compatible_copy(source, target):
    before = load_article(source)
    missing = _missing_image_relationships(source)
    main_document_rids = missing.get("word/_rels/document.xml.rels", set())
    with tempfile.TemporaryDirectory(prefix="baijiahao-docx-") as temporary:
        working_source = Path(source)
        if missing:
            working_source = Path(temporary) / Path(source).name
            _write_without_missing_images(source, working_source, missing)
        original = Document(working_source)
        output = Document()
        for paragraph in original.paragraphs:
            new = output.add_paragraph()
            try:
                style_name = getattr(paragraph.style, "name", None)
                if style_name:
                    new.style = style_name
            except (AttributeError, KeyError):
                pass
            new.alignment = paragraph.alignment
            for run in paragraph.runs:
                copied = new.add_run(run.text)
                copied.bold, copied.italic, copied.underline = run.bold, run.italic, run.underline
                for blip in run._r.xpath('.//a:blip'):
                    rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if not rid:
                        continue
                    try:
                        raw = original.part.related_parts[rid].blob
                    except KeyError:
                        if rid in main_document_rids:
                            continue
                        raise
                    with Image.open(io.BytesIO(raw)) as image:
                        if image.format not in {'JPEG', 'PNG', 'GIF', 'BMP', 'TIFF'}:
                            buffer = io.BytesIO()
                            image.save(buffer, format='PNG')
                            raw = buffer.getvalue()
                    copied.add_picture(io.BytesIO(raw), width=Inches(6))
        output.save(target)
    after = load_article(target)
    if before.text != after.text or [pixels(i.data) for i in before.images] != [pixels(i.data) for i in after.images]:
        raise ValueError('兼容导入副本的正文或图片核对不一致，停止上传')
    return target
